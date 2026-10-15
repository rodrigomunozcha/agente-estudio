"""
Reuniones y entrevistas: lo que sale de una conversacion grabada que no es clase.

Por que existe. El estudiante tambien graba reuniones de grupo y entrevistas de
trabajo para poder volver a ellas, y de esas conversaciones salen cosas que
hacer. Se agrego al mismo agente y no a uno aparte (decision del 30-09-2026)
porque casi todo ya estaba: el AirDrop, Whisper, el bot, la secretaria y los
avisos. Lo unico nuevo es el analisis.

Como entra. "REUNIÓN" es un ramo mas en ramos_adicionales de config.json, asi
que aparece en la ventana de confirmacion y el estudiante lo elige a mano, como
elige un ramo. finalizar_clase manda aqui en vez de a la skill de clases.

Que deja:
  - una nota en la boveda, en Reuniones/,
  - un .docx en "2 Areas/Reuniones post Claude/" de OneDrive,
  - un resumen por Telegram: una linea de que fue, las acciones con su fecha y los
    seguimientos, con el .docx adjunto,
  - las acciones con fecha, anotadas para que la secretaria las recuerde a las 9:00,
  - los seguimientos: lo que la otra parte prometio con plazo ("te respondemos en una
    semana"), anotados como pendientes de tipo seguimiento para el dia que vencen.

Solo se afirma lo que tiene cita en la transcripcion, y eso lo comprueba Python
(evaluaciones.cita_respaldada), no el modelo: una accion o una promesa inventada hace
quedar mal a quien actua sobre ella. Una accion sin cita queda en la nota marcada para
verificar, pero no va a Telegram ni a los recordatorios. Un seguimiento sin cita se
descarta.

Nombres de personas: aqui SI van, como excepcion pedida por el estudiante el
30-09-2026 (en una entrevista importa quien dijo que). Correos, telefonos, RUT y
datos bancarios siguen sin ir: la boveda no los admite (ver su Inicio.md).

El analisis es una llamada al modelo que no puede leer ni escribir archivos:
recibe la transcripcion en el prompt y devuelve un JSON, que Python valida y
escribe. Puede usar la busqueda web, solo para la parte de "que investigar".

A mano, sobre una transcripcion ya hecha (por ejemplo, una que quedo en
Output/Sin clasificar):
    python3 -m orquestador.reuniones <transcripcion.txt> <AAAA-MM-DD> "<nombre>"
"""
import json
import re
import urllib.parse
from datetime import date, datetime, timedelta
from pathlib import Path

from claude_agent_sdk import AssistantMessage, ClaudeAgentOptions, ResultMessage, TextBlock

from . import evaluaciones
from .archivado import archivar_audio
from .config import cargar_config, dir_pendientes
from .ensayo import es_ensayo
from .hoja_html import HERRAMIENTAS_PROHIBIDAS
from .nombres import nombre_base, sanitizar_nombre_archivo
from .notificaciones import notificar_evento, notificar_reunion
from .skill_runner import CLI_PATH, PROJECT_ROOT, consultar, describir_error_sdk
from .uso import registrar_uso

RAMO = "REUNIÓN"
CARPETA_BOVEDA = "Reuniones"
CARPETA_DOCX = Path.home() / "Library/CloudStorage/OneDrive-Personal/2 Areas/Reuniones post Claude"
RUTA_INSTRUCCIONES = Path(__file__).parent / "plantilla_reunion" / "instrucciones.md"
MARCA_INICIO, MARCA_FIN = "===REUNION_INICIO===", "===REUNION_FIN==="
MAX_TURNS = 12  # la busqueda web gasta turnos, este es el seguro
# El modelo mas capaz del plan Pro (decision del 14-10-2026, issue #1): las reuniones
# son pocas y su calidad pesa mas que su tiempo. Sin costo extra: va por la sesion Pro.
MODELO = "claude-opus-5-5"
# Un plazo relativo mas largo que esto no es un seguimiento que valga recordar: es
# casi seguro un numero mal escuchado.
MAX_PLAZO_DIAS = 180
TIMEOUT_SEGUNDOS = 20 * 60
PROHIBIDAS = [h for h in HERRAMIENTAS_PROHIBIDAS if h != "WebSearch"]

_ENCABEZADO_SIN_CLASIFICAR = re.compile(r"\A.*?\n-{20,}\n\n", re.S)


# --------------------------------------------------------------------------
# El analisis
# --------------------------------------------------------------------------

def extraer(texto: str) -> dict | None:
    """El JSON entre las marcas, con lo minimo que el resto necesita. None si
    no vino o no sirve: mejor avisar que guardar una reunion a medias."""
    inicio, fin = texto.find(MARCA_INICIO), texto.rfind(MARCA_FIN)
    if inicio == -1 or fin <= inicio:
        return None
    crudo = re.sub(r"^```(?:json)?\s*|\s*```$", "", texto[inicio + len(MARCA_INICIO):fin].strip())
    try:
        datos = json.loads(crudo)
    except ValueError:
        return None
    if not isinstance(datos, dict) or not str(datos.get("titulo", "")).strip() or not datos.get("resumen"):
        return None
    for lista in ("decisiones", "acciones", "seguimientos", "preguntas_abiertas", "investigar"):
        if not isinstance(datos.get(lista), list):
            datos[lista] = []
    for lista in ("acciones", "seguimientos"):
        datos[lista] = [a for a in datos[lista] if isinstance(a, dict) and a.get("que")]
    if not isinstance(datos.get("entrevista"), dict):
        datos["entrevista"] = None
    return datos


async def analizar(texto: str, fecha: str, nombre: str, slug: str) -> tuple[dict | None, str | None]:
    options = ClaudeAgentOptions(
        cwd=str(PROJECT_ROOT),
        setting_sources=[],
        system_prompt=RUTA_INSTRUCCIONES.read_text(encoding="utf-8"),
        allowed_tools=["WebSearch"],
        disallowed_tools=PROHIBIDAS,
        permission_mode="bypassPermissions",
        cli_path=str(CLI_PATH),
        model=MODELO,
        max_turns=MAX_TURNS,
    )
    prompt = (f"Fecha de la conversación: {fecha}\nNombre que le puso el estudiante a la grabación: "
              f"{nombre}\n\n===TRANSCRIPCIÓN===\n{texto}\n===FIN===\n\nAnalízala siguiendo tus "
              f"instrucciones y devuelve solo el JSON entre {MARCA_INICIO} y {MARCA_FIN}.")
    partes, final, error = [], None, None
    async for mensaje in consultar(prompt=prompt, options=options):
        if isinstance(mensaje, AssistantMessage):
            partes += [b.text for b in mensaje.content if isinstance(b, TextBlock)]
        elif isinstance(mensaje, ResultMessage):
            registrar_uso("reunion", slug, mensaje)
            final = mensaje.result
            if mensaje.is_error:
                error = describir_error_sdk(mensaje)
    datos = extraer(final or "") or extraer("\n".join(partes))
    if datos is None:
        return None, (f"hubo un problema de conexión ({error})" if error else "el modelo no devolvió el análisis")
    return datos, None


# --------------------------------------------------------------------------
# Lo que arma Python
# --------------------------------------------------------------------------

def _linea_accion(a: dict) -> str:
    cuando = a.get("fecha") or a.get("cuando") or "sin fecha"
    extra = [x for x in (a.get("hora"), a.get("lugar")) if x]
    return (f"{a['que']} - {a.get('quien') or 'sin responsable claro'}, {cuando}"
            + (f" ({', '.join(extra)})" if extra else ""))


def _fecha_seguimiento(s: dict, fecha_reunion: str) -> str | None:
    """La fecha en que vence una promesa. La calcula Python, como la proxima clase de
    un pedido: la que se dijo, o la de la reunion mas el plazo. None si no hay como."""
    fecha = evaluaciones.validar_fecha(s.get("fecha"))
    plazo = s.get("plazo_dias")
    if fecha is None and isinstance(plazo, int) and not isinstance(plazo, bool) and 0 < plazo <= MAX_PLAZO_DIAS:
        fecha = (date.fromisoformat(fecha_reunion) + timedelta(days=plazo)).isoformat()
    return fecha if fecha and fecha > fecha_reunion else None


def validar_con_transcripcion(datos: dict, transcripcion: str, fecha_reunion: str) -> dict:
    """Marca cada accion con "respaldada" segun su cita este o no en la transcripcion,
    y deja solo los seguimientos respaldados, con su fecha de vencimiento calculada."""
    for a in datos["acciones"]:
        a["respaldada"] = evaluaciones.cita_respaldada(a.get("cita") or "", transcripcion)
    datos["seguimientos"] = [dict(s, fecha=_fecha_seguimiento(s, fecha_reunion))
                             for s in datos.get("seguimientos") or []
                             if evaluaciones.cita_respaldada(s.get("cita") or "", transcripcion)]
    return datos


def secciones(datos: dict) -> list[tuple[str, list[str] | str]]:
    """El contenido en orden, una sola vez, para la nota y para el .docx."""
    s: list[tuple[str, list[str] | str]] = [("Resumen", datos["resumen"])]
    if datos.get("calidad_audio"):
        s.append(("Calidad del audio", datos["calidad_audio"]))
    if datos["decisiones"]:
        s.append(("Decisiones", datos["decisiones"]))
    if datos["acciones"]:
        s.append(("Acciones", [_linea_accion(a) + (f". Se dijo: \"{a['cita']}\"" if a.get("cita") else "")
                               + (" (no encontré esta frase en la grabación: verifícala)"
                                  if a.get("respaldada") is False else "")
                               for a in datos["acciones"]]))
    if datos.get("seguimientos"):
        s.append(("Seguimientos (lo que prometió la otra parte)", [
            f"{x['que']} - {x.get('quien') or 'la otra parte'}, {x.get('fecha') or x.get('cuando') or 'sin fecha'}"
            + (f". Se dijo: \"{x['cita']}\"" if x.get("cita") else "") for x in datos["seguimientos"]]))
    if datos["preguntas_abiertas"]:
        s.append(("Preguntas abiertas", datos["preguntas_abiertas"]))
    e = datos.get("entrevista")
    if e:
        s.append(("La entrevista", [
            f"Pregunta: {p.get('pregunta', '')}\nTu respuesta: {p.get('tu_respuesta', '')}\n"
            f"Sugerencia para mejorarla: {p.get('como_mejorarla', '')}"
            for p in e.get("preguntas") or [] if isinstance(p, dict)]))
        if e.get("impresion_general"):
            s.append(("Cómo te fue", e["impresion_general"]))
        if e.get("proximos_pasos"):
            s.append(("Próximos pasos", e["proximos_pasos"]))
    for t in datos["investigar"]:
        if isinstance(t, dict) and t.get("tema"):
            fuentes = ", ".join(t.get("fuentes") or []) or "conocimiento general"
            s.append((f"Para investigar: {t['tema']}",
                      f"{t.get('por_que', '')}\n\n{t.get('explicacion', '')}\n\nFuente: {fuentes}".strip()))
    return s


def nota_markdown(datos: dict, fecha: str, nombre: str) -> str:
    partes = ["---", f"fecha: {fecha}", f"tipo: {datos.get('tipo') or 'reunion'}", "tags:", "  - reunion",
              "---", "", f"# {datos['titulo']}", "",
              f"Grabación: {nombre}. Transcrita por Whisper y analizada por el agente de estudio.", ""]
    for titulo, cuerpo in secciones(datos):
        partes.append(f"## {titulo}\n")
        if isinstance(cuerpo, list):
            partes += [f"- {x}".replace("\n", "\n  ") for x in cuerpo] + [""]
        else:
            partes += [cuerpo, ""]
    return "\n".join(partes)


def escribir_docx(datos: dict, fecha: str, nombre: str, destino: Path) -> Path:
    from docx import Document  # python-docx, ya instalado (requirements.txt)
    doc = Document()
    doc.add_heading(datos["titulo"], level=0)
    doc.add_paragraph(f"{fecha} · Grabación: {nombre}")
    for titulo, cuerpo in secciones(datos):
        doc.add_heading(titulo, level=1)
        for linea in (cuerpo if isinstance(cuerpo, list) else [cuerpo]):
            doc.add_paragraph(linea, style="List Bullet" if isinstance(cuerpo, list) else None)
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)
    return destino


def _libre(ruta: Path) -> Path:
    """Nunca pisar: si el nombre esta ocupado, "nombre (2)"."""
    destino, n = ruta, 2
    while destino.exists():
        destino, n = ruta.with_name(f"{ruta.stem} ({n}){ruta.suffix}"), n + 1
    return destino


def enlace_calendario(accion: dict, titulo_reunion: str) -> str | None:
    """
    El enlace que abre Google Calendar con el evento ya lleno. Formato
    action=TEMPLATE, que Google no documenta oficialmente pero que varias fuentes
    independientes describen igual (verificado el 30-09-2026): dates en
    AAAAMMDDTHHMMSS/AAAAMMDDTHHMMSS con ctz para la zona horaria, o AAAAMMDD/AAAAMMDD
    (el segundo dia exclusivo) para uno de dia completo. Sin fecha no hay enlace.

    Color: siempre Mandarina, pedido del estudiante el 30-09-2026. Va con
    pprop=eventColor:6, tambien sin documentacion oficial, y 6 es Mandarina en la
    numeracion de colores de Google (sus clases usan 9, Arandano, y sus
    evaluaciones 11, Tomate). El lugar va solo si hay que ir fisicamente a alguna
    parte: una reunion en linea o una tarea no llevan lugar.
    """
    fecha = evaluaciones.validar_fecha(accion.get("fecha"))
    if not fecha:
        return None
    inicio = datetime.fromisoformat(fecha)
    hora = re.fullmatch(r"(\d{1,2}):(\d{2})", str(accion.get("hora") or ""))
    if hora:
        inicio = inicio.replace(hour=int(hora.group(1)), minute=int(hora.group(2)))
        fechas = f"{inicio:%Y%m%dT%H%M%S}/{inicio + timedelta(hours=1):%Y%m%dT%H%M%S}"
    else:
        fechas = f"{inicio:%Y%m%d}/{inicio + timedelta(days=1):%Y%m%d}"
    detalle = f"De: {titulo_reunion}" + (f"\nSe dijo: \"{accion['cita']}\"" if accion.get("cita") else "")
    if accion.get("presencial") is True:
        detalle += "\nHay que moverse: agrega un aviso 2 horas antes."
    params = {"action": "TEMPLATE", "text": accion["que"], "dates": fechas, "ctz": "America/Santiago",
              "details": detalle, "pprop": "eventColor:6"}
    if accion.get("presencial") is True and accion.get("lugar"):
        params["location"] = accion["lugar"]
    return "https://calendar.google.com/calendar/render?" + urllib.parse.urlencode(params, quote_via=urllib.parse.quote)


def anotar_acciones(datos: dict, fecha: str, ensayo: bool) -> int:
    """Las acciones respaldadas con fecha futura y los seguimientos con fecha, para el
    mensaje de las 9:00 y /pendientes. Las anota evaluaciones.anotar_pendiente, que es
    el unico que escribe el archivo."""
    por_anotar = [(a["que"] + (f" ({a['quien']})" if a.get("quien") else ""), f, a, "")
                  for a in datos["acciones"]
                  if a.get("respaldada") and (f := evaluaciones.validar_fecha(a.get("fecha"))) and f >= fecha]
    por_anotar += [(s["que"] + (f" ({s['quien']})" if s.get("quien") else ""), s["fecha"], s, "seguimiento")
                   for s in datos.get("seguimientos") or [] if s.get("fecha")]
    n = 0
    for que, f, a, tipo in por_anotar:
        try:
            n += evaluaciones.anotar_pendiente(RAMO, f, que, "reunion", textual=a.get("cita") or "",
                                               clase=fecha, reunion=datos["titulo"], tipo=tipo, ensayo=ensayo)
        except (OSError, ValueError, evaluaciones.CandadoOcupado):
            return n  # la reunion se guarda igual: solo faltan sus recordatorios
    return n


def _dd_mm(iso: str | None) -> str:
    return f"{iso[8:10]}-{iso[5:7]}" if iso else "sin fecha"


def resumen_telegram(datos: dict, anotadas: int) -> str:
    """Se lee en el telefono, caminando: una linea de que fue, las acciones con su fecha
    y los seguimientos. Solo lo respaldado por una cita; el detalle queda en el .docx."""
    bloques = [f"📝 {datos.get('una_linea') or datos['titulo']}"]
    acciones = [a for a in datos["acciones"] if a.get("respaldada")]
    if acciones:
        bloques.append("✅ Acciones\n" + "\n".join(
            f"• {a['que']} - {a.get('quien') or 'sin responsable claro'}, "
            + (_dd_mm(evaluaciones.validar_fecha(a.get("fecha"))) if a.get("fecha") else a.get("cuando") or "sin fecha")
            for a in acciones))
    if datos.get("seguimientos"):
        bloques.append("⏳ Seguimientos\n" + "\n".join(
            f"• {s['que']} - {s.get('quien') or 'la otra parte'}, {_dd_mm(s.get('fecha'))}"
            for s in datos["seguimientos"]))
    temas = [t["tema"] for t in datos["investigar"] if isinstance(t, dict) and t.get("tema")]
    if temas:
        bloques.append("🔎 En el .docx te explico: " + ", ".join(temas))
    if anotadas:
        bloques.append(f"⏰ Te recuerdo {anotadas} a las 9:00 del día que toca")
    return "\n\n".join(bloques)


def guardar(datos: dict, fecha: str, nombre: str, config: dict, transcripcion: str) -> tuple[Path, Path]:
    """Escribe la nota y el .docx, anota las acciones y los seguimientos, y avisa.
    Devuelve (nota, docx). La transcripcion es para comprobar cada cita aqui, justo
    antes de afirmarla, tambien en un analisis guardado de una corrida anterior."""
    validar_con_transcripcion(datos, transcripcion, fecha)
    base = sanitizar_nombre_archivo(f"{fecha} - {datos['titulo']}")
    carpeta_nota = Path(config["rutas"]["vault_obsidian"]) / CARPETA_BOVEDA
    carpeta_nota.mkdir(parents=True, exist_ok=True)
    nota = _libre(carpeta_nota / f"Reunión - {base}.md")
    nota.write_text(nota_markdown(datos, fecha, nombre), encoding="utf-8")
    carpeta_docx = Path(config["rutas"]["output"]) / RAMO if es_ensayo(config) else CARPETA_DOCX
    # La nota de Obsidian conserva su nombre: renombrar notas rompe enlaces. El
    # .docx se llama como el audio y el .txt (ver nombres.nombre_base).
    docx = escribir_docx(datos, fecha, nombre, _libre(carpeta_docx / f"{nombre_base(fecha, datos['titulo'], RAMO)}.docx"))
    anotadas = anotar_acciones(datos, fecha, es_ensayo(config))
    notificar_reunion(datos["titulo"], resumen_telegram(datos, anotadas), docx)
    if not es_ensayo(config):
        for a in datos["acciones"]:
            url = (enlace_calendario(a, datos["titulo"])
                   if a.get("evento") is True and a.get("respaldada") else None)
            if url:
                cuando = a["fecha"] + (f", {a['hora']}" if a.get("hora") else "")
                notificar_evento(a["que"], cuando, url,
                                 a["lugar"] if a.get("presencial") is True and a.get("lugar") else "")
    return nota, docx


async def procesar_reunion(trabajo_metadata: dict, config: dict, bitacora=None) -> Path:
    """Lo que finalizar_clase llama en vez de la skill cuando el ramo es REUNIÓN."""
    slug = trabajo_metadata.get("slug") or trabajo_metadata["fecha"]
    guardado = dir_pendientes() / f"{slug}_reunion.json"
    nombre = ", ".join(Path(a).name for a in trabajo_metadata["archivos_originales"])
    texto = Path(trabajo_metadata["archivo_texto"]).read_text(encoding="utf-8")
    if guardado.exists():
        # Si una corrida anterior llego al analisis y se corto despues, no se paga de nuevo.
        datos = json.loads(guardado.read_text(encoding="utf-8"))
    else:
        datos, motivo = await analizar(texto, trabajo_metadata["fecha"], nombre, slug)
        if datos is None:
            raise RuntimeError(f"No se pudo analizar la reunión: {motivo}. La transcripción quedó guardada.")
        guardado.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    nota, docx = guardar(datos, trabajo_metadata["fecha"], nombre, config, texto)
    if bitacora is not None:
        bitacora.archivo_creado(nota)
    trabajo = dict(trabajo_metadata, archivos=trabajo_metadata["archivos_originales"])
    archivar_audio(trabajo, datos["titulo"], config, bitacora)
    return docx


if __name__ == "__main__":
    import sys

    import anyio

    if len(sys.argv) != 4:
        print('Uso: python3 -m orquestador.reuniones <transcripcion.txt> <AAAA-MM-DD> "<nombre>"')
        raise SystemExit(1)
    _texto = _ENCABEZADO_SIN_CLASIFICAR.sub("", Path(sys.argv[1]).read_text(encoding="utf-8"), count=1)
    _fecha = date.fromisoformat(sys.argv[2]).isoformat()

    async def _una():
        with anyio.fail_after(TIMEOUT_SEGUNDOS):
            return await analizar(_texto, _fecha, sys.argv[3], f"manual {_fecha}")

    _datos, _motivo = anyio.run(_una)
    if _datos is None:
        print(f"No se pudo analizar: {_motivo}")
        raise SystemExit(1)
    _nota, _docx = guardar(_datos, _fecha, sys.argv[3], cargar_config(), _texto)
    print(f"  nota -> {_nota}\n  docx -> {_docx}")
